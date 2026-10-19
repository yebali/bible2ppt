"""자막 슬라이드 배치(긴 절 나누기), 템플릿으로 PPT 만들기, 미리보기 그림 그리기.

템플릿 규칙: 첫 슬라이드에 {본문} 글자가 든 글상자 하나. 글상자 채우기가 곧 자막 띠이고,
나머지 모양(배경·색·글꼴·위치)은 그대로 복제된다.
"""
import colorsys
import copy
import io
import re
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path

from lxml import etree
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import nsuri, qn
from pptx.util import Emu, Inches, Pt

BODY = "{본문}"
DEFAULT_LINES = 2  # 기본 한 장 최대 줄 수 = 기본 템플릿 띠 높이(한글 줄 수)
SAFETY = 0.96  # PowerPoint와 글자 폭 계산이 조금 달라도 줄이 늘어나지 않게 여유를 둠
# 글꼴 이름 -> 폭 측정용 파일 후보. 없으면 FALLBACK(맑은 고딕, 맥에서는 애플 SD 고딕)으로 잼
# ponytail: 템플릿을 맑은 고딕이 아닌 글꼴로 바꿔 줄 나눔이 어긋나면 여기에 글꼴 파일을 추가 (예: "나눔고딕": ["NanumGothic.ttf"])
FONT_FILES = {}
BOLD_FILES = {"malgun.ttf": "malgunbd.ttf"}
FALLBACK = ["malgun.ttf", "AppleSDGothicNeo.ttc"]
_SHAPES = {qn(t) for t in ("p:sp", "p:pic", "p:grpSp", "p:graphicFrame", "p:cxnSp")}
_R = f"{{{nsuri('r')}}}"


@dataclass
class Slide:
    book: str                                   # "왕상"
    verse: str                                  # "19:2"
    piece: str = ""                             # 긴 절이 나뉘면 "a", "b"...
    lines: list = field(default_factory=list)   # 첫 줄은 "왕상 19:2a. ..."로 시작

    @property
    def ref(self):
        return f"{self.book} {self.verse}{self.piece}"

    @property
    def text(self):  # 장절 표시를 뺀 본문
        return " ".join(self.lines)[len(self.ref) + 2:]


@lru_cache(maxsize=256)
def _font(name, bold, px):
    for f in FONT_FILES.get(name, []) + FALLBACK:
        try:
            return ImageFont.truetype(BOLD_FILES.get(f, f) if bold else f, max(1, px))
        except OSError:
            continue
    return ImageFont.load_default(max(1, px))


@lru_cache(maxsize=65536)
def _measure(name, bold, px, s):
    """글자 폭 (px). 줄 나누기가 같은 줄을 여러 번 재므로 문자열별로 기억."""
    return _font(name, bold, px).getlength(s)


@dataclass
class Style:
    face: str
    size: float     # pt
    bold: bool
    color: tuple
    spacing: float  # 줄 간격 배수
    align: str = "l"  # 문단 정렬: l 왼쪽, ctr 가운데, r 오른쪽

    def font(self, scale):
        """pt -> px 배율에 맞춘 글꼴."""
        return _font(self.face, self.bold, round(self.size * scale))

    def width(self, s):
        return _measure(self.face, self.bold, round(self.size * 4), s) / 4

    @property
    def line_h(self):
        asc, desc = self.font(4).getmetrics()
        return (asc + desc) / 4 * self.spacing


def _scheme(slide):
    """테마 색 표 {이름: 'RRGGBB'} (accent1, dk1 ... 와 bg1·tx1 같은 별칭)."""
    master = slide.slide_layout.slide_master
    theme = etree.fromstring(master.part.part_related_by(RT.THEME).blob)
    colors = {}
    for c in theme.find(f".//{qn('a:clrScheme')}"):
        v = c[0]
        colors[etree.QName(c).localname] = v.get("val") if etree.QName(v).localname == "srgbClr" else v.get("lastClr")
    cmap = master._element.find(qn("p:clrMap"))
    for alias, name in (cmap.attrib.items() if cmap is not None else ()):
        colors[alias] = colors.get(name)
    return colors


def _rgba(fill, scheme):
    """a:solidFill -> (r, g, b, a). 직접 고른 색, 테마 색, 밝기 조절(lumMod/lumOff), 투명도까지 계산."""
    for clr in [] if fill is None else fill:
        kind = etree.QName(clr).localname
        rgb = {"srgbClr": clr.get("val"), "schemeClr": scheme.get(clr.get("val")),
               "sysClr": clr.get("lastClr")}.get(kind)
        if not rgb:
            continue
        r, g, b = (int(rgb[i:i + 2], 16) / 255 for i in (0, 2, 4))
        mods = {etree.QName(m).localname: int(m.get("val")) / 100000 for m in clr}
        if "lumMod" in mods or "lumOff" in mods:  # PowerPoint 색 고르기의 '더 밝게/어둡게'
            h, lum, sat = colorsys.rgb_to_hls(r, g, b)
            r, g, b = colorsys.hls_to_rgb(h, min(1, max(0, lum * mods.get("lumMod", 1) + mods.get("lumOff", 0))), sat)
        return round(r * 255), round(g * 255), round(b * 255), round(255 * mods.get("alpha", 1))
    return None


def _style(p, scheme, default_color=(255, 255, 255)):
    r = p.find(qn("a:r"))
    rpr = r.find(qn("a:rPr")) if r is not None else None
    rpr = rpr if rpr is not None else etree.Element(qn("a:rPr"))
    face = next((el.get("typeface") for tag in ("a:ea", "a:latin")
                 if (el := rpr.find(qn(tag))) is not None), "맑은 고딕")
    ppr = p.find(qn("a:pPr"))
    pct = ppr.find(f"{qn('a:lnSpc')}/{qn('a:spcPct')}") if ppr is not None else None
    return Style(face, int(rpr.get("sz", 1800)) / 100, rpr.get("b") in ("1", "true"),
                 (_rgba(rpr.find(qn("a:solidFill")), scheme) or default_color)[:3],
                 int(pct.get("val")) / 100000 if pct is not None else 1.0,
                 ppr.get("algn", "l") if ppr is not None else "l")


def _find_box(slide):
    return next((s for s in slide.shapes if s.has_text_frame and BODY in s.text_frame.text), None)


def _body_para(box):
    return next(p for p in box.text_frame.paragraphs if BODY in p.text)


class Template:
    def __init__(self, path):
        self.path = Path(path)
        prs = Presentation(self.path)
        if not len(prs.slides):
            raise ValueError("템플릿에 슬라이드가 없어요")
        slide = prs.slides[0]
        box = _find_box(slide)
        if box is None:
            raise ValueError(f"템플릿 첫 슬라이드에 {BODY} 글상자가 없어요")
        self.slide_w, self.slide_h = prs.slide_width.pt, prs.slide_height.pt
        self.left, self.top, self.width, self.height = box.left.pt, box.top.pt, box.width.pt, box.height.pt
        tf = box.text_frame
        self.ins = [tf.margin_left.pt, tf.margin_top.pt, tf.margin_right.pt, tf.margin_bottom.pt]
        self.anchor = tf._txBody.bodyPr.get("anchor", "t")  # 띠 안에서 글자 세로 위치: t 위, ctr 가운데, b 아래
        scheme = _scheme(slide)
        self.body = _style(_body_para(box)._p, scheme)
        # 배경: 슬라이드에 없으면 레이아웃, 마스터 순으로 (단색만)
        self.bg = next((c for owner in (slide, slide.slide_layout, slide.slide_layout.slide_master)
                        if (bg := owner._element.cSld.bg) is not None
                        and (c := _rgba(bg.find(f"{qn('p:bgPr')}/{qn('a:solidFill')}"), scheme))), (0, 255, 0, 255))
        self.fill = _rgba(box._element.spPr.find(qn("a:solidFill")), scheme) or (0, 0, 0, 0)
        self.extras = []  # 띠 말고 템플릿에 있는 그림·글상자 (미리보기에 같이 그림): ((x, y, w, h), 그림|글, 글 서식)
        for sh in slide.shapes:
            if sh.shape_id == box.shape_id or sh.width is None:
                continue
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                geo = (sh.left.pt, sh.top.pt, sh.width.pt, sh.height.pt)
                self.extras.append((geo, Image.open(io.BytesIO(sh.image.blob)).convert("RGBA"), None))
            elif sh.has_text_frame and sh.text_frame.text.strip():
                t = sh.text_frame  # 글자는 글상자 여백만큼 안쪽에서 시작
                geo = (sh.left.pt + t.margin_left.pt, sh.top.pt + t.margin_top.pt, sh.width.pt, sh.height.pt)
                self.extras.append((geo, t.text, _style(t.paragraphs[0]._p, scheme, (0, 0, 0))))

    @property
    def text_w(self):
        """띠(글상자) 안쪽 폭."""
        return self.width - self.ins[0] - self.ins[2]

    def fits(self, s):
        return self.body.width(s) <= self.text_w * SAFETY

    @property
    def room(self):
        """띠(글상자) 안쪽 높이. 띠 크기는 템플릿 그대로 고정."""
        return self.height - self.ins[1] - self.ins[3]

    @property
    def fit_lines(self):
        """띠에 들어가는 한글 줄 수."""
        return int((self.room + 0.5) // self.body.line_h)


def _pieces(words, label, fits, max_lines):
    """어절을 줄에 채우고, max_lines를 넘으면 다음 조각으로. 조각마다 첫 줄 맨 앞에 label(n)."""
    pieces, i = [], 0
    while i < len(words):
        start, lines, line = i, [], label(len(pieces))
        while i < len(words):
            cand = f"{line} {words[i]}".lstrip()
            if fits(cand) or not line:
                line, i = cand, i + 1
            elif len(lines) + 1 < max_lines:
                # ponytail: 한 줄보다 긴 어절 하나는 자르지 않음 (PowerPoint가 알아서 줄바꿈)
                lines.append(line)
                line, i = words[i], i + 1
            else:
                break
        if i == start:  # 줄 수 1에 첫 어절도 안 들어가는 경우: 억지로 넣어 무한 반복 방지
            line, i = f"{line} {words[i]}".lstrip(), i + 1
        pieces.append(lines + [line])
    return pieces


SENTENCE_END = re.compile(r"[.?!][\"'”’)]*$|[\"”’]$")
PAUSE = re.compile(r"[,;:][\"'”’)]*$")
CLAUSE_END = re.compile(r"(되|니|매|고|며|여|어|서|나|면|라|다|요|늘|사)$")  # 개역 문체의 이어 주는 말 끝 (이르되, 하니, 하며…)


def _break_cost(word):
    """이 어절 뒤에서 다음 장으로 넘길 때의 벌점: 문장 끝 0, 쉼표 1, 이어 주는 말 2, 그 밖 6."""
    return 0 if SENTENCE_END.search(word) else 1 if PAUSE.search(word) else 2 if CLAUSE_END.search(word) else 6


def split(ref, text, fits, max_lines, width):
    """한 절 -> [(표시 장절, 줄들)]. 나뉘면 '왕상 19:2a.', '왕상 19:2b.' ...
    장 수는 최소로 두고, 끊기 좋은 자리(문장 끝·쉼표·이어 주는 말)를 우선하면서 장끼리 분량을 고르게 나눔."""
    words = text.split()
    if not words:
        return [(ref, [f"{ref}."])]
    whole = _pieces(words, lambda n: f"{ref}.", fits, max_lines)  # 짧은 표시로 먼저: 한 장에 들어가면 안 나눔
    if len(whole) == 1:
        return [(ref, whole[0])]
    label = lambda k: f"{ref}{chr(97 + k)}."
    n = len(_pieces(words, label, fits, max_lines))  # 필요한 최소 장 수 (앞에서부터 채운 결과도 아래 후보에 들어 있음)

    @lru_cache(maxsize=None)
    def lay(k, i, j):  # words[i:j]가 k번째 장 한 장에 들어가면 줄들, 아니면 None
        p = _pieces(words[i:j], lambda _: label(k), fits, max_lines)
        return p[0] if len(p) == 1 else None

    acc = [0.0]
    for w in words:
        acc.append(acc[-1] + width(w + " "))
    avg = acc[-1] / n
    uneven = lambda i, j: ((acc[j] - acc[i] - avg) / avg) ** 2 * 10  # 장 분량이 평균에서 벗어난 만큼 벌점

    @lru_cache(maxsize=None)
    def best(k, i):  # words[i:]를 k번째 장부터 나눌 때 (벌점, 끊는 위치들)
        if k == n - 1:
            return (uneven(i, len(words)), ()) if lay(k, i, len(words)) else (float("inf"), ())
        out = (float("inf"), ())
        for j in range(i + 1, len(words)):
            if not lay(k, i, j):
                break
            cost, rest = best(k + 1, j)
            out = min(out, (cost + uneven(i, j) + _break_cost(words[j - 1]), (j,) + rest))
        return out

    cuts = (0,) + best(0, 0)[1] + (len(words),)
    return [(label(k)[:-1], lay(k, cuts[k], cuts[k + 1])) for k in range(n)]


def layout(items, tpl, max_lines):
    """[(책 약어, 장:절, 본문)] -> [Slide]. 한 장 줄 수는 설정값과 템플릿 띠에 들어가는 줄 수 중 작은 쪽."""
    per = min(max_lines, tpl.fit_lines)
    if per < 1:
        raise ValueError("템플릿 띠가 너무 낮아서 한 줄도 안 들어가요. 템플릿에서 띠(글상자)를 키워 주세요")
    slides = []
    for book, verse, body in items:
        ref = f"{book} {verse}"
        parts = split(ref, body, tpl.fits, per, tpl.body.width)
        slides += [Slide(book, verse, r[len(ref):], lines) for r, lines in parts]  # 나뉜 조각만 r 끝에 a, b…
    return slides


# ───────── PPT 저장 ─────────

def _insert(p, el):
    end = p.find(qn("a:endParaRPr"))
    end.addprevious(el) if end is not None else p.append(el)


def _set_lines(para, lines):
    """문단의 첫 글자 서식을 유지한 채 줄들로 바꿈 (줄 사이는 줄바꿈 <a:br/>)."""
    proto = copy.deepcopy(para.runs[0]._r)
    para.clear()
    p = para._p
    rpr = proto.find(qn("a:rPr"))
    for n, line in enumerate(lines):
        if n:
            br = etree.Element(qn("a:br"))
            if rpr is not None:
                br.append(copy.deepcopy(rpr))
            _insert(p, br)
        r = copy.deepcopy(proto)
        r.find(qn("a:t")).text = line
        _insert(p, r)


def _remap(el, rids):
    for node in el.iter():
        for k, v in node.attrib.items():
            if k.startswith(_R) and v in rids:
                node.set(k, rids[v])
    return el


def _clone(prs, src):
    """src 슬라이드를 그대로 복제 (그림 등 연결 관계 포함)."""
    dst = prs.slides.add_slide(src.slide_layout)
    tree = dst.shapes._spTree
    for el in list(tree):
        if el.tag in _SHAPES:
            tree.remove(el)  # 레이아웃 자리표시자 제거
    rids = {}
    for rid, rel in src.part.rels.items():
        if not rel.reltype.endswith(("/slideLayout", "/notesSlide")):
            target = rel.target_ref if rel.is_external else rel.target_part
            rids[rid] = dst.part.relate_to(target, rel.reltype, rel.is_external)
    ext = tree.find(qn("p:extLst"))
    for el in src.shapes._spTree:
        if el.tag in _SHAPES:
            new = _remap(copy.deepcopy(el), rids)
            ext.addprevious(new) if ext is not None else tree.append(new)
    bg = src._element.cSld.bg
    if bg is not None:
        dst._element.cSld.insert(0, _remap(copy.deepcopy(bg), rids))
    return dst


def build(slides, tpl):
    """템플릿을 복제해 자막을 채운 Presentation (저장은 따로)."""
    prs = Presentation(tpl.path)
    src = prs.slides[0]
    lst = prs.slides._sldIdLst
    originals = list(lst)  # 템플릿에 원래 있던 슬라이드 (첫 장 말고 더 있어도 전부 결과에서 뺌)
    for s in slides:
        box = _find_box(_clone(prs, src))  # 띠 위치·크기는 템플릿 그대로, 글자만 바꿈
        _set_lines(_body_para(box), s.lines)
    for sld in originals:
        lst.remove(sld)
        prs.part.drop_rel(sld.rId)
    return prs


def make_default_template(path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(0x00, 0xFF, 0x00)
    m, ins = Inches(0.12), Inches(0.08)
    # 띠 높이: 맑은 고딕 44pt 한글 DEFAULT_LINES줄 + 여백. 맥에서 만든 템플릿을 윈도우로 옮겨도 들어가게 줄 높이에 여유(1.35배)
    line_h = max(Style("맑은 고딕", 44, False, (0, 0, 0), 1.0).line_h, 44 * 1.35)
    h = Emu(Pt(DEFAULT_LINES * line_h + 6) + 2 * ins)
    box = slide.shapes.add_textbox(m, prs.slide_height - m - h, prs.slide_width - 2 * m, h)
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(0, 0, 0)
    etree.SubElement(box._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr")), qn("a:alpha"), val="60000")
    box.line.fill.background()
    tf = box.text_frame
    tf.word_wrap, tf.auto_size, tf.vertical_anchor = True, MSO_AUTO_SIZE.NONE, MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = Inches(0.2)
    tf.margin_top = tf.margin_bottom = ins
    r = tf.paragraphs[0].add_run()
    r.text = BODY
    r.font.size, r.font.name = Pt(44), "맑은 고딕"
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    r._r.get_or_add_rPr().find(qn("a:latin")).addnext(etree.Element(qn("a:ea"), typeface="맑은 고딕"))  # 한글 글꼴
    prs.save(path)


# ───────── 미리보기 그림 ─────────

@lru_cache(maxsize=8)
def _backdrop(tpl, width_px):
    """본문 글자를 뺀 슬라이드 (배경, 그림·글상자, 띠). 장마다 같아서 크기별로 한 번만 그림."""
    k = width_px / tpl.slide_w  # pt -> px
    img = Image.new("RGBA", (width_px, round(tpl.slide_h * k)), tpl.bg)
    for (x, y, w, h), what, st in tpl.extras:
        if st is None:  # 그림
            pic = what.resize((max(1, round(w * k)), max(1, round(h * k))))
            img.paste(pic, (round(x * k), round(y * k)), pic)
        else:  # 글상자
            ImageDraw.Draw(img).text((x * k, y * k), what, font=st.font(k), fill=st.color)
    band = Image.new("RGBA", img.size, (0, 0, 0, 0))
    ImageDraw.Draw(band).rectangle([tpl.left * k, tpl.top * k, (tpl.left + tpl.width) * k,
                                    (tpl.top + tpl.height) * k], fill=tpl.fill)
    return Image.alpha_composite(img, band)


def render(slide, tpl, width_px):
    k = width_px / tpl.slide_w  # pt -> px
    img = _backdrop(tpl, width_px).copy()
    d = ImageDraw.Draw(img)
    st = tpl.body
    spare = tpl.room - len(slide.lines) * st.line_h  # 띠 안 남는 높이
    y = (tpl.top + tpl.ins[1] + {"ctr": spare / 2, "b": spare}.get(tpl.anchor, 0)) * k
    f = st.font(k)
    room_w = tpl.text_w * k
    for line in slide.lines:
        spare_w = room_w - f.getlength(line)  # 가운데·오른쪽 정렬이면 그만큼 밀어서
        x = (tpl.left + tpl.ins[0]) * k + {"ctr": spare_w / 2, "r": spare_w}.get(st.align, 0)
        d.text((x, y), line, font=f, fill=st.color)
        y += st.line_h * k
    return img.convert("RGB")
